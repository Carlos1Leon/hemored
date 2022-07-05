# -*- coding: utf-8 -*-

import io
import os
import datetime
import tempfile

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from lxml import etree

from odoo import fields, api, models

from ..models import constants

MN_Hemovigilancia = '.'.join([constants.MODULO, 'hemovigilancia'])


class Hemovigilancia(models.TransientModel):
    _name = MN_Hemovigilancia

    def es_diresa(self):
        return self.env.user.has_group('hemored.group_coordinador_diris_diresa_hemored')

    def _default_diresa(self):
        if self.es_diresa():
            return self.env.user.diresa_id.id
        return False

    banco_id = fields.Many2many(
        'hemored.banco_sangre',
        'hemovigilancia_banco_sangre_rel',
        'hemovigilancia_id',
        'banco_id',
        'Centro de Hemoterapia Banco de Sangre'
    )
    diresa_id = fields.Many2many(
        'renipress.diresa',
        'hemovigilancia_diresa_rel',
        'hemovigilancia_id',
        'diresa_id',
        'DIRIS/DIRESA/GERESA',
        default=lambda self: self._default_diresa()
    )
    anho_id = fields.Many2many(
        'minsa.anho',
        'hemovigilancia_anho_rel',
        'hemovigilancia_id',
        'anho_id',
        'Año'
    )
    periodo_id = fields.Many2many(
        'minsa.periodo',
        'hemovigilancia_periodo_rel',
        'hemovigilancia_id',
        'periodo_id',
        'Periodo'
    )
    tipo_banco = fields.Selection(
        '_selection_tipo_banco',
        'Tipo de CHBS',
    )
    institucion = fields.Selection(
        constants.SELECTION_INSTITUCION,
        'Institución'
    )
    estado_validacion = fields.Selection(
        constants.SELECTION_ESTADO_FICHA,
        'Estado validación',
    )
    fecha = fields.Date(
        string='Fecha',
        default=fields.Date.context_today
    )

    @api.model
    def fields_view_get(self, view_id=None, view_type=False, toolbar=False, submenu=False):
        res = super(Hemovigilancia, self).fields_view_get(view_id=view_id, view_type=view_type, toolbar=toolbar, submenu=submenu)
        es_diresa = self.env.user.has_group('hemored.group_coordinador_diris_diresa_hemored')
        if view_type == 'form':
            doc = etree.XML(res['arch'])
            my_attrs = '{"readonly": true}'
            if es_diresa:
                for node in doc.xpath("//field[@name='diresa_id']"):
                    node.set('modifiers', my_attrs)
            res['arch'] = etree.tostring(doc)
        return res

    @api.model
    def _selection_tipo_banco(self):
        return self.env['basecatalogo.catalogo'].get_catalogo('_selection_tipo_banco')

    @api.onchange('diresa_id', 'tipo_banco', 'institucion')
    def onchange_diresa_id(self):
        domain = {}
        if self.diresa_id and self.tipo_banco and self.institucion:
            domain = {'banco_id': '''[('diresa_id', 'in', %s), ('tipo_banco', '=', %s), ('institucion', '=', %s)]''' % (
                str(self.diresa_id.ids), self.tipo_banco, self.institucion)}
        elif self.diresa_id and self.tipo_banco:
            domain = {'banco_id': '''[('diresa_id', 'in', %s), ('tipo_banco', '=', %s)]''' % (
                str(self.diresa_id.ids), self.tipo_banco)}
        elif self.diresa_id and self.institucion:
            domain = {'banco_id': '''[('diresa_id', 'in', %s), ('institucion', '=', %s)]''' % (
                str(self.diresa_id.ids), self.institucion)}
        elif self.diresa_id:
            domain = {'banco_id': '''[('diresa_id', 'in', %s)]''' % str(self.diresa_id.ids)}
        elif self.tipo_banco:
            domain = {'banco_id': '''[('tipo_banco', '=', %s)]''' % (self.tipo_banco)}
        elif self.institucion:
            domain = {'banco_id': '''[('institucion', '=', %s)]''' % (self.institucion)}
        else:
            domain = {}
        self.banco_id = False
        return {'domain': domain}

    @api.multi
    def print_excel(self):
        self.ensure_one()
        date = datetime.datetime.now() - datetime.timedelta(hours=5)
        url = '/web/binary/download_export_hemovigilancia?id={}&filename=Reporte_Hemovigilancia_{}.xlsx'.format(self.id, date.strftime('%d-%m-%Y %H:%M:%S'))
        value = {
            'type': 'ir.actions.act_url',
            'url': url,
            'target': 'new'
        }
        return value

    @api.multi
    def action_export_excel(self):
        self.ensure_one()
        archivo = 'FORMATO_BASE_REPORTE_HEMOVIGILANCIA.xlsx'
        ruta = '{}/../documents/{}'.format(os.path.dirname(os.path.abspath(__file__)), archivo)
        book = load_workbook(ruta)

        fobj = tempfile.NamedTemporaryFile(suffix='.xlsx', dir='/tmp/', delete=False)
        fname = fobj.name
        fobj.close()

        w_sheet, w_sheet2, w_sheet3 = book.get_sheet_by_name('Tamizaje Unidades'), book.get_sheet_by_name('Paquete Globular'), book.get_sheet_by_name('Eliminación Glóbulos Rojos')
        w_sheet4, w_sheet5 = book.get_sheet_by_name('Total Tamizaje Unidades'), book.get_sheet_by_name('Gráficos Tamizaje Unidades')

        constrains, constrains_4 = [], []
        constrains_2, constrains_3 = '', ''

        if len(self.banco_id) > 0:
            constrains_1 = str(self.banco_id[0].id)
            for i in range(len(self.banco_id)):
                constrains_1 += ', ' + str(self.banco_id[i].id)
            constrains.append('fe.banco_id in ({})'.format(constrains_1))

        if len(self.diresa_id) > 0:
            constrains_1 = str(self.diresa_id[0].id)
            for i in range(len(self.diresa_id)):
                constrains_1 += ', ' + str(self.diresa_id[i].id)
            constrains.append('fe.diresa_id in ({})'.format(constrains_1))

        if len(self.anho_id) > 0:
            constrains_1 = str(self.anho_id[0].id)
            for i in range(len(self.anho_id)):
                constrains_1 += ', ' + str(self.anho_id[i].id)
            constrains.append('fe.anho_id in ({})'.format(constrains_1))
            constrains_4.append('fe.anho_id in ({})'.format(constrains_1))

        if len(self.periodo_id) > 0:
            constrains_1 = str(self.periodo_id[0].id)
            for i in range(len(self.periodo_id)):
                constrains_1 += ', ' + str(self.periodo_id[i].id)
            constrains.append('fe.periodo_id in ({})'.format(constrains_1))
            constrains_4.append('fe.periodo_id in ({})'.format(constrains_1))
            constrains_2 = 'mp.name as periodo, '
            constrains_3 = ', mp.name'

        if self.tipo_banco:
            constrains.append('fe.tipo_banco = \'{}\''.format(self.tipo_banco))

        if self.institucion:
            constrains.append('fe.institucion = \'{}\''.format(self.institucion))

        if self.estado_validacion:
            constrains.append('fe.state = \'{}\''.format(self.estado_validacion))

        constrains = ' AND '.join(constrains) or 'true'
        constrains_4 = ' AND '.join(constrains_4) or 'true'

        map_field_celda = constants.MAP_FIELD_CELDA
        fields = [
            'SUM({field}) as {field}'.format(field=field)
            for field in map_field_celda.keys()
        ]
        fields = ', '.join(fields)

        sql = '''
                SELECT re.name as chbs, re.codigo_eess, rd.name as diresa, fe.institucion, ma.name as anho, fe.tipo_banco, {}{}
                FROM hemored_ficha_estadistica fe
                    INNER JOIN hemored_banco_sangre bs on fe.banco_id = bs.id
                    INNER JOIN renipress_eess re on bs.renipress_id = re.id
                    INNER JOIN renipress_diresa rd on re.diresa_id = rd.id
                    INNER JOIN minsa_anho ma on fe.anho_id = ma.id
                    INNER JOIN minsa_periodo mp on fe.periodo_id = mp.id
                WHERE {}
                GROUP BY re.name, re.codigo_eess, rd.name, fe.institucion, ma.name, fe.tipo_banco{};
        '''.format(constrains_2, fields, constrains, constrains_3)

        self._cr.execute(sql)

        temp, temp2, tipo_banco, institucion = 5, 3, '', ''
        total = self._cr.dictfetchall()
        for i in range(len(total)):
            w_sheet['A' + str(temp + i)], w_sheet2['A' + str(temp2 + i)], w_sheet3['A' + str(temp2 + i)] = i + 1, i + 1, i + 1
            for field, valor in total[i].items():
                if valor:
                    if field == 'chbs':
                        w_sheet['B' + str(temp + i)], w_sheet2['B' + str(temp2 + i)], w_sheet3['B' + str(temp2 + i)] = valor, valor, valor
                    elif field == 'codigo_eess':
                        w_sheet['C' + str(temp + i)], w_sheet2['C' + str(temp2 + i)], w_sheet3['C' + str(temp2 + i)] = valor, valor, valor
                    elif field == 'tipo_banco':
                        if valor == '1':
                            tipo_banco = constants.TIPO_BANCO_1
                        elif valor == '2':
                            tipo_banco = constants.TIPO_BANCO_2
                        w_sheet['D' + str(temp + i)], w_sheet2['D' + str(temp2 + i)], w_sheet3['D' + str(temp2 + i)] = tipo_banco, tipo_banco, tipo_banco
                    elif field == 'diresa':
                        w_sheet['E' + str(temp + i)], w_sheet2['E' + str(temp2 + i)], w_sheet3['E' + str(temp2 + i)] = valor, valor, valor
                    elif field == 'institucion':
                        if valor == '1':
                            institucion = constants.SELECTION_INSTITUCION_1
                        elif valor == '2':
                            institucion = constants.SELECTION_INSTITUCION_2
                        elif valor == '3':
                            institucion = constants.SELECTION_INSTITUCION_3
                        elif valor == '4':
                            institucion = constants.SELECTION_INSTITUCION_4
                        elif valor == '5':
                            institucion = constants.SELECTION_INSTITUCION_5
                        elif valor == '6':
                            institucion = constants.SELECTION_INSTITUCION_6
                        elif valor == '7':
                            institucion = constants.SELECTION_INSTITUCION_7
                        elif valor == '8':
                            institucion = constants.SELECTION_INSTITUCION_8
                        elif valor == '9':
                            institucion = constants.SELECTION_INSTITUCION_9
                        elif valor == '10':
                            institucion = constants.SELECTION_INSTITUCION_10
                        elif valor == '11':
                            institucion = constants.SELECTION_INSTITUCION_11
                        w_sheet['F' + str(temp + i)], w_sheet2['F' + str(temp2 + i)], w_sheet3['F' + str(temp2 + i)] = institucion, institucion, institucion
                    elif field == 'anho':
                        w_sheet['G' + str(temp + i)], w_sheet2['G' + str(temp2 + i)], w_sheet3['G' + str(temp2 + i)] = valor, valor, valor
                    elif field == 'periodo':
                        w_sheet['H' + str(temp + i)], w_sheet2['H' + str(temp2 + i)], w_sheet3['H' + str(temp2 + i)] = valor, valor, valor
                    elif field == 'cantidad_ds_hgr':
                        w_sheet2[map_field_celda[field] + str(temp2 + i)] = valor
                    elif field == 'cantidad_da_hgr':
                        w_sheet2[map_field_celda[field] + str(temp2 + i)] = valor
                    elif field == 'cantidad_cv_hgr':
                        w_sheet3[map_field_celda[field] + str(temp2 + i)] = valor
                    elif field == 'cantidad_cmitt_hgr':
                        w_sheet3[map_field_celda[field] + str(temp2 + i)] = valor
                    elif field == 'cantidad_ca_hgr':
                        w_sheet3[map_field_celda[field] + str(temp2 + i)] = valor
                    elif field == 'cantidad_ct_hgr':
                        w_sheet3[map_field_celda[field] + str(temp2 + i)] = valor
                    elif field == 'cantidad_cp_hgr':
                        w_sheet3[map_field_celda[field] + str(temp2 + i)] = valor
                    elif field == 'cantidad_co_hgr':
                        w_sheet3[map_field_celda[field] + str(temp2 + i)] = valor
                    else:
                        w_sheet[map_field_celda[field] + str(temp + i)] = valor
                else:
                    if field == 'periodo':
                        w_sheet['H' + str(temp + i)], w_sheet2['H' + str(temp2 + i)], w_sheet3['H' + str(temp2 + i)] = '', '', ''
                    elif field == 'cantidad_ds_hgr':
                        w_sheet2[map_field_celda[field] + str(temp2 + i)] = 0
                    elif field == 'cantidad_da_hgr':
                        w_sheet2[map_field_celda[field] + str(temp2 + i)] = 0
                    elif field == 'cantidad_cv_hgr':
                        w_sheet3[map_field_celda[field] + str(temp2 + i)] = 0
                    elif field == 'cantidad_cmitt_hgr':
                        w_sheet3[map_field_celda[field] + str(temp2 + i)] = 0
                    elif field == 'cantidad_ca_hgr':
                        w_sheet3[map_field_celda[field] + str(temp2 + i)] = 0
                    elif field == 'cantidad_ct_hgr':
                        w_sheet3[map_field_celda[field] + str(temp2 + i)] = 0
                    elif field == 'cantidad_cp_hgr':
                        w_sheet3[map_field_celda[field] + str(temp2 + i)] = 0
                    elif field == 'cantidad_co_hgr':
                        w_sheet3[map_field_celda[field] + str(temp2 + i)] = 0
                    else:
                        w_sheet[map_field_celda[field] + str(temp + i)] = 0

        map_field_celda = constants.MAP_FIELD_CELDA_2
        fields = [
            'SUM({field}) as {field}'.format(field=field)
            for field in map_field_celda.keys()
        ]
        fields = ', '.join(fields)

        sql = '''
                SELECT ma.name as anho, {}{}
                FROM hemored_ficha_estadistica fe
                    INNER JOIN hemored_banco_sangre bs on fe.banco_id = bs.id
                    INNER JOIN renipress_eess re on bs.renipress_id = re.id
                    INNER JOIN renipress_diresa rd on re.diresa_id = rd.id
                    INNER JOIN minsa_anho ma on fe.anho_id = ma.id
                    INNER JOIN minsa_periodo mp on fe.periodo_id = mp.id
                WHERE {}
                GROUP BY ma.name{};
        '''.format(constrains_2, fields, constrains_4, constrains_3)

        self._cr.execute(sql)

        temp2 = 3
        total = self._cr.dictfetchall()
        for i in range(len(total)):
            w_sheet4['A' + str(temp2 + i)] = i + 1
            for field, valor in total[i].items():
                if valor:
                    if field == 'anho':
                        w_sheet4['B' + str(temp2 + i)] = valor
                    elif field == 'periodo':
                        w_sheet4['C' + str(temp2 + i)] = valor
                    else:
                        w_sheet4[map_field_celda[field] + str(temp2 + i)] = valor
                else:
                    if field == 'anho':
                        w_sheet4['B' + str(temp2 + i)] = 0
                    elif field == 'periodo':
                        w_sheet4['C' + str(temp2 + i)] = 0
                    else:
                        w_sheet4[map_field_celda[field] + str(temp2 + i)] = 0
        barchart_vih = BarChart()
        data_vih = Reference(w_sheet4, min_col=4, max_col=6, min_row=2, max_row=len(total) + 2)
        cat_vih = Reference(w_sheet4, min_col=3, max_col=3, min_row=3, max_row=len(total) + 2)
        barchart_vih.add_data(data_vih, titles_from_data=True)
        barchart_vih.set_categories(cat_vih)
        w_sheet5.add_chart(barchart_vih, 'A1')
        barchart_vih.title = 'VIH'
        barchart_vih.style = 2

        barchart_hbsag = BarChart()
        data_hbsag = Reference(w_sheet4, min_col=7, max_col=9, min_row=2, max_row=len(total) + 2)
        cat_hbsag = Reference(w_sheet4, min_col=3, max_col=3, min_row=3, max_row=len(total) + 2)
        barchart_hbsag.add_data(data_hbsag, titles_from_data=True)
        barchart_hbsag.set_categories(cat_hbsag)
        w_sheet5.add_chart(barchart_hbsag, 'H1')
        barchart_hbsag.title = 'HBsAg'
        barchart_hbsag.style = 2

        barchart_hepc = BarChart()
        data_hepc = Reference(w_sheet4, min_col=10, max_col=12, min_row=2, max_row=len(total) + 2)
        cat_hepc = Reference(w_sheet4, min_col=3, max_col=3, min_row=3, max_row=len(total) + 2)
        barchart_hepc.add_data(data_hepc, titles_from_data=True)
        barchart_hepc.set_categories(cat_hepc)
        w_sheet5.add_chart(barchart_hepc, 'O1')
        barchart_hepc.title = 'Hep C'
        barchart_hepc.style = 2

        barchart_antihbc = BarChart()
        data_antihbc = Reference(w_sheet4, min_col=13, max_col=15, min_row=2, max_row=len(total) + 2)
        cat_antihbc = Reference(w_sheet4, min_col=3, max_col=3, min_row=3, max_row=len(total) + 2)
        barchart_antihbc.add_data(data_antihbc, titles_from_data=True)
        barchart_antihbc.set_categories(cat_antihbc)
        w_sheet5.add_chart(barchart_antihbc, 'A16')
        barchart_antihbc.title = 'Anti-HBc'
        barchart_antihbc.style = 2

        barchart_htlv = BarChart()
        data_htlv = Reference(w_sheet4, min_col=16, max_col=18, min_row=2, max_row=len(total) + 2)
        cat_htlv = Reference(w_sheet4, min_col=3, max_col=3, min_row=3, max_row=len(total) + 2)
        barchart_htlv.add_data(data_htlv, titles_from_data=True)
        barchart_htlv.set_categories(cat_htlv)
        w_sheet5.add_chart(barchart_htlv, 'H16')
        barchart_htlv.title = 'HTLV I/II'
        barchart_htlv.style = 2

        barchart_sifilis = BarChart()
        data_sifilis = Reference(w_sheet4, min_col=19, max_col=21, min_row=2, max_row=len(total) + 2)
        cat_sifilis = Reference(w_sheet4, min_col=3, max_col=3, min_row=3, max_row=len(total) + 2)
        barchart_sifilis.add_data(data_sifilis, titles_from_data=True)
        barchart_sifilis.set_categories(cat_sifilis)
        w_sheet5.add_chart(barchart_sifilis, 'O16')
        barchart_sifilis.title = 'Sifilis'
        barchart_sifilis.style = 2

        barchart_chagas = BarChart()
        data_chagas = Reference(w_sheet4, min_col=22, max_col=24, min_row=2, max_row=len(total) + 2)
        cat_chagas = Reference(w_sheet4, min_col=3, max_col=3, min_row=3, max_row=len(total) + 2)
        barchart_chagas.add_data(data_chagas, titles_from_data=True)
        barchart_chagas.set_categories(cat_chagas)
        w_sheet5.add_chart(barchart_chagas, 'A32')
        barchart_chagas.title = 'Chagas'
        barchart_chagas.style = 2

        barchart_otros = BarChart()
        data_otros = Reference(w_sheet4, min_col=25, max_col=27, min_row=2, max_row=len(total) + 2)
        cat_otros = Reference(w_sheet4, min_col=3, max_col=3, min_row=3, max_row=len(total) + 2)
        barchart_otros.add_data(data_otros, titles_from_data=True)
        barchart_otros.set_categories(cat_otros)
        w_sheet5.add_chart(barchart_otros, 'H32')
        barchart_otros.title = 'Otros'
        barchart_otros.style = 2

        fname = io.BytesIO()
        book.save(fname)
        fname.seek(0)
        data = fname.read()
        fname.close()
        return data
